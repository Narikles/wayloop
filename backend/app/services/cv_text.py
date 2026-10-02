"""Extraction du texte d'un CV (PDF, DOCX, TXT). Pas d'OCR : un scan est signalé comme illisible."""
from __future__ import annotations

import io
import re

ALLOWED_MIME = {
    "application/pdf": ".pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": ".docx",
    "text/plain": ".txt",
}


def guess_suffix(filename: str, mime: str | None) -> str:
    ext = ("." + filename.rsplit(".", 1)[-1].lower()) if "." in filename else ""
    if ext in {".pdf", ".docx", ".txt"}:
        return ext
    return ALLOWED_MIME.get(mime or "", "")


def extract_text(data: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        return _pdf(data)
    if suffix == ".docx":
        return _docx(data)
    if suffix == ".txt":
        for enc in ("utf-8", "cp1252", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
    return ""


def _pdf(data: bytes) -> str:
    try:
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(p.extract_text(x_tolerance=1.5, y_tolerance=3) or "") for p in pdf.pages[:10]]
        text = "\n".join(pages)
    except Exception:  # noqa: BLE001 - PDF malformé : on tente pypdf
        text = ""
    if len(text.strip()) < 50:
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((p.extract_text() or "") for p in reader.pages[:10])
        except Exception:  # noqa: BLE001
            pass
    return _clean(text)


def _docx(data: bytes) -> str:
    try:
        import docx

        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return _clean("\n".join(parts))
    except Exception:  # noqa: BLE001
        return ""


def _clean(text: str) -> str:
    text = text.replace("\x00", "").replace("­", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def is_readable(text: str) -> bool:
    return len(re.findall(r"[A-Za-zÀ-ÿ]{3,}", text)) >= 25
